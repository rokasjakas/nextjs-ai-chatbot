"""Converts the team's Google Sheets calendar (xlsx) into the website's
Renginiai / Paėmimai JSON import file.

Usage: python3 scripts/sheets_to_events_json.py calendar.xlsx renginiai-importas.json
Then open Renginiai → „⬆ Importuoti“ on the website and choose the JSON file.
"""
import datetime, json, re, sys, warnings, collections
import openpyxl

warnings.filterwarnings("ignore")
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

TIME_LABELS = ['Sandėlyje','Išvažiuoti','Lokacijoje','Soundcheck','Repeticija','Paruošti iki','Pradžia','Pabaiga','Demontavimas']
GROUP_COLS = {11:'Daiktų rūšiavimas', 12:'Montavimas diena prieš', 14:'Montavimas renginio dieną', 16:'Demontavimas',
              18:'Garsas', 19:'VJ', 20:'Šviesa', 21:'Budi renginyje'}
NOTE_FILLS = {'FFCFE2F3','FFC9DAF8','FFD0E0E3'}
POS_RE = re.compile(r'^((?:[AVL]\d*(?:/(?:OP|CREW))?)|rigg|DeRigg\??|Elektr\.?)[\s/]+(.+)$', re.I)
DAYTIME_RE = re.compile(r'^\s*\d{1,2}\s*d\b|^\s*\d{1,2}:\d{2}|\d{1,2}:\d{2}\s*$', re.I)
WORK_RE = re.compile(r'rūšiav|rusiav|sandėlio darb|sandelio darb|servis', re.I)

def text(v):
    if v is None: return ''
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d') if v.time()==datetime.time(0) else v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, datetime.time): return v.strftime('%H:%M')
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def fill(cell):
    try:
        if cell.fill and cell.fill.fill_type:
            rgb = cell.fill.fgColor.rgb
            return rgb if isinstance(rgb, str) else None
    except Exception: pass
    return None

def parse_date(v):
    """-> (date, dateEnd) from a datetime or text like 2026-09-04/06 or 2026-01-09/11."""
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d'), ''
    s = text(v)
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})\s*(?:/\s*(\d{1,2}))?', s)
    if not m: return None, ''
    d = f'{m.group(1)}-{m.group(2)}-{m.group(3)}'
    end = ''
    if m.group(4):
        day = int(m.group(4)); y, mo, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if day < dd: mo += 1
        if mo > 12: mo, y = 1, y+1
        try: end = datetime.date(y, mo, day).isoformat()
        except ValueError: end = ''
    return d, end

people = collections.OrderedDict()
UP, LO = 'A-ZĄČĘĖĮŠŲŪŽ', 'a-ząčęėįšųūž'
NAME_RE = re.compile(rf'^([{UP}][{LO}]+ [{UP}][{LO}]+(?:-[{UP}][{LO}]+)?)(?:\s|$)')
NEED_CATS = [('LED ekranai', r'led|ekran|\bp\d|p2|p3'), ('Garsas', r'gars|audio|mikrof|kolonėl'), ('Šviesa', r'švies|svies|apšviet|light'),
             ('Video', r'video|projekt|hdmi|kamer'), ('Scena', r'scen|pakyl'), ('Santvaros', r'ferm|santvar|truss|gravity'), ('Elektra', r'elektr|generator|63a|32a')]
needs = collections.OrderedDict()

def person_entry(raw, fillrgb):
    raw = re.sub(r'\s+', ' ', raw).strip()
    m = POS_RE.match(raw)
    pos, name = (m.group(1), m.group(2).strip()) if m else ('', raw)
    status = 'gali' if fillrgb == 'FFFF9900' else 'negali' if fillrgb == 'FFFF0000' else 'siulomas' if fillrgb == 'FFB4A7D6' else ''
    m2 = NAME_RE.match(name)   # "Vitas Račyla su savo auto" -> "Vitas Račyla"
    if m2: people[m2.group(1)] = people.get(m2.group(1), 0) + 1
    return {'pos': pos, 'person': name, 'status': status}

def convert_events(ws, tag):
    out = []
    starts = [r for r in range(2, ws.max_row+1) if ws.cell(r,1).value not in (None, '')]
    for idx, r0 in enumerate(starts):
        date, date_end = parse_date(ws.cell(r0,1).value)
        if not date: continue       # month header rows (SAUSIS'26 …)
        r1 = (starts[idx+1] if idx+1 < len(starts) else ws.max_row+1) - 1
        rows = range(r0, r1+1)
        title = text(ws.cell(r0,2).value)
        labels = [text(ws.cell(r,9).value) for r in rows]
        has_times = any(l for l in labels)
        # rows without a timeline and a venue are notes / warehouse work, not full events
        is_work = (not has_times) and (not text(ws.cell(r0,6).value))
        # crew groups
        groups = []
        for col, gtitle in GROUP_COLS.items():
            notes, entries = [], []
            for r in rows:
                c = ws.cell(r, col); v = text(c.value)
                if not v or v in ('q',): continue
                f = fill(c)
                if f in NOTE_FILLS or (not entries and DAYTIME_RE.search(v) and not POS_RE.match(v)):
                    notes.append(re.sub(r'\s+', ' ', v))
                else:
                    entries.append(person_entry(v, f))
            if notes or entries or not is_work:
                groups.append({'id': f'g{col}', 'title': gtitle, 'note': ' / '.join(notes), 'entries': entries})
        # logistics: V = there, X = back ("IVECO MFY418 - Algis, Valaitis")
        logistics = []
        for col, day in ((22, ''), (24, 'atgal')):
            for r in rows:
                v = text(ws.cell(r, col).value)
                if not v: continue
                veh, _, drv = v.partition(' - ')
                item = {'day': day, 'vehicle': veh.strip(), 'driver': drv.strip()}
                if day == 'atgal' and any(l['vehicle'].lower()==item['vehicle'].lower() and l['driver']==item['driver'] for l in logistics): continue
                logistics.append(item)
        notes = '\n'.join(text(ws.cell(r,25).value) for r in rows if text(ws.cell(r,25).value))
        ev_id = f'imp-{tag}-r{r0}'
        if is_work:
            crew_entries = [e for g in groups for e in g['entries']]
            time = next((g['note'] for g in groups if g['note']), '')
            out.append({'id': ev_id, 'kind': 'work', 'date': date, 'title': title or 'Sandėlio darbai', 'time': time,
                        'crew': [{'id':'g1','title':'Žmonės','note':'','entries':crew_entries}],
                        'logistics': logistics, 'notes': notes, 'createdAt': f'{date}T00:00:{r0%60:02d}', 'imported': True})
            continue
        # times: "Sandėlyje:" + J; "Paruošti iki:9d" keeps the suffix
        times = {l: '' for l in TIME_LABELS}
        extra = []
        for r in rows:
            lab = text(ws.cell(r,9).value)
            if not lab: continue
            name, _, suffix = lab.partition(':')
            name = name.strip()
            t = text(ws.cell(r,10).value)
            val = (suffix.strip()+' '+t).strip() if suffix.strip() else t
            if name in times: times[name] = val
            else: extra.append({'label': name, 'time': val})
        lines = [l for l in title.split('\n') if l.strip() and not set(l.strip()) <= set('-')]
        need = re.sub(r'\s+', ' ', text(ws.cell(r0,3).value)).strip()
        ev_needs = [cat for cat, rx in NEED_CATS if need and re.search(rx, need, re.I)]
        for n in ev_needs: needs.setdefault(n, None)
        mgr = re.sub(r'\s*\n\s*', ' · ', text(ws.cell(r0,8).value)).strip(' ·')
        out.append({
            'id': ev_id, 'kind': 'event', 'date': date, 'dateEnd': date_end,
            'name': lines[0] if lines else '(be pavadinimo)', 'artist': ' '.join(lines[1:]),
            'needs': ev_needs, 'needsNote': need, 'location': re.sub(r'\s*\n\s*', ', ', text(ws.cell(r0,6).value)),
            'organizer': re.sub(r'\s*\n\s*', ', ', text(ws.cell(r0,7).value)), 'manager': mgr, 'projectId': None,
            'times': [{'label': l, 'time': times[l]} for l in TIME_LABELS] + extra,
            'crew': groups, 'logistics': logistics, 'notes': notes,
            'createdAt': f'{date}T00:00:{r0%60:02d}', 'imported': True,
        })
    return out

events = []
for name, tag in (('Renginiai 2026','2026'), ('Renginiai 2027','2027')):
    if name in wb.sheetnames: events += convert_events(wb[name], tag)

# ---------- rentals ----------
STATUS = {'gražinta':('grazinta',''), 'grazinta':('grazinta',''), 'paimta':('paimta',''), 'reikia paimti':('reikia',''),
          'nereikia gražinti':('grazinta','Nereikia grąžinti'), 'nepatvirtinta':('reikia','Nepatvirtinta'),
          'reikia gražinti':('paimta','Reikia grąžinti'), 'dalis daiktu grazinta':('paimta','Dalis daiktų grąžinta'),
          'reikia nuvežti':('reikia','Reikia nuvežti')}
ws = wb['Įrangos nuoma paėmimasgražinima']
rentals, suppliers = [], collections.OrderedDict()
today = datetime.date.today().isoformat()
def split_dt(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d'), ('' if v.time()==datetime.time(0) else v.strftime('%H:%M')), ''
    s = text(v)
    return '', '', s
for r in range(2, ws.max_row+1):
    items = text(ws.cell(r,2).value)
    if not items: continue
    raw_status = text(ws.cell(r,1).value)
    st, note = STATUS.get(raw_status.lower(), (None, raw_status if raw_status and raw_status!='Pasirinkti' else ''))
    pd, pt, ptxt = split_dt(ws.cell(r,5).value)
    rd, rt, rtxt = split_dt(ws.cell(r,6).value)
    if st is None:
        ref = rd or pd
        st = 'grazinta' if (ref and ref < today) or not ref else 'reikia'
    place_raw = re.sub(r'\s+', ' ', text(ws.cell(r,3).value))
    place, _, address = place_raw.partition(',')
    if not address:
        m = re.match(r'^(.*?)\s+([A-ZĄČĘĖĮŠŲŪŽ][\wąčęėįšųūž]+\s+g\.?\s*\d.*)$', place_raw)
        if m: place, address = m.group(1), m.group(2)
    place, address = place.strip(), address.strip()
    if place:
        s_ = suppliers.setdefault(place, {'n':0, 'address':''})
        s_['n'] += 1
        if address and not s_['address']: s_['address'] = address
    notes = [n for n in [note, ('Paėmimas: '+ptxt) if ptxt else '', ('Grąžinimas: '+rtxt) if rtxt else '', text(ws.cell(r,7).value)] if n]
    rentals.append({'id': f'imp-rt-r{r}', 'status': st, 'items': items, 'place': place, 'address': address,
                    'contact': text(ws.cell(r,4).value), 'pickupDate': pd, 'pickupTime': pt, 'returnDate': rd, 'returnTime': rt,
                    'notes': '\n'.join(notes), 'person': text(ws.cell(r,8).value), 'eventId': None,
                    'createdAt': f'2020-01-01T00:00:00.{r:06d}', 'imported': True})

data = {'format': 'es-events-import', 'version': 1, 'createdAt': datetime.datetime.now().isoformat(timespec='seconds'),
        'events': events, 'rentals': rentals,
        'people': [k for k, n in people.items() if n >= 2], 'needs': list(needs.keys()),
        'suppliers': [{'name': k, 'address': v['address'], 'contact': ''} for k, v in suppliers.items()
                      if v['n'] >= 2 and not re.search(r'\d|\bg\.', k) and 3 <= len(k) <= 40]}
json.dump(data, open(sys.argv[2], 'w', encoding='utf-8'), ensure_ascii=False)
kinds = collections.Counter(e['kind'] for e in events)
print('events', len(events), dict(kinds), '| months', sorted(collections.Counter(e['date'][:7] for e in events).items())[:3], '…')
print('rentals', len(rentals), dict(collections.Counter(r['status'] for r in rentals)))
print('people', len(data['people']), '| needs', data['needs'], '| suppliers', len(data['suppliers']))
